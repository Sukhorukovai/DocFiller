# -*- coding: utf-8 -*-
"""
Операции над .docx: runs, замены, поиск блоков.

Все функции работают с объектом `docx.Document` (python-docx).
Здесь нет ни OCR, ни парсинга — только работа с XML-представлением.
"""
import copy
from datetime import date

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .core.text import norm_text
from .logging_setup import log


# =====================================================================
# rPr / runs
# =====================================================================

def make_plain_rPr(src_rPr=None):
    new_rPr = OxmlElement('w:rPr')
    if src_rPr is not None:
        for tag in ('w:rFonts', 'w:sz', 'w:szCs', 'w:lang', 'w:kern'):
            el = src_rPr.find(qn(tag))
            if el is not None:
                new_rPr.append(copy.deepcopy(el))
    return new_rPr


def set_run_plain(r_el, text, src_rPr=None, color_hex="FF0000"):
    for ch in list(r_el):
        if ch.tag in (qn('w:t'), qn('w:tab'), qn('w:br'), qn('w:cr')):
            r_el.remove(ch)

    rPr = r_el.find(qn('w:rPr'))
    if rPr is not None:
        r_el.remove(rPr)
    rPr = make_plain_rPr(src_rPr)
    r_el.insert(0, rPr)

    tokens = (text or '').split('\n') if False else None  # noqa
    import re as _re
    for tok in _re.split(r'([\t\n])', text or ''):
        if not tok:
            continue
        if tok == '\t':
            r_el.append(OxmlElement('w:tab'))
        elif tok == '\n':
            r_el.append(OxmlElement('w:br'))
        else:
            t = OxmlElement('w:t')
            t.text = tok
            t.set(qn('xml:space'), 'preserve')
            r_el.append(t)

    if color_hex:
        rrPr = r_el.find(qn('w:rPr'))
        c = rrPr.find(qn('w:color'))
        if c is not None:
            rrPr.remove(c)
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color_hex)
        rrPr.append(c)


def run_text(r_el):
    parts = []
    for ch in r_el:
        if ch.tag == qn('w:t'):
            parts.append(ch.text or '')
        elif ch.tag == qn('w:tab'):
            parts.append('\t')
        elif ch.tag in (qn('w:br'), qn('w:cr')):
            parts.append('\n')
    return ''.join(parts)


def iter_runs_ordered(elem):
    result = []
    for child in elem:
        if child.tag == qn('w:r'):
            result.append((child, run_text(child)))
        elif child.tag == qn('w:hyperlink'):
            result.extend(iter_runs_ordered(child))
    return result


# =====================================================================
# Обход документа
# =====================================================================

def iter_unique_cells(row):
    seen = set()
    result = []
    for c in row.cells:
        tc = c._tc
        if id(tc) in seen:
            continue
        seen.add(id(tc))
        result.append(c)
    return result


def iter_all_paragraphs(doc):
    seen = set()

    def walk_paras(paras):
        for p in paras:
            if id(p._element) in seen:
                continue
            seen.add(id(p._element))
            yield p

    yield from walk_paras(doc.paragraphs)

    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for c in iter_unique_cells(row):
                    yield from walk_paras(c.paragraphs)
                    yield from walk_tables(c.tables)

    yield from walk_tables(doc.tables)

    for section in doc.sections:
        for hf in (section.header, section.footer):
            yield from walk_paras(hf.paragraphs)
            yield from walk_tables(hf.tables)


# =====================================================================
# Замена диапазона в абзаце
# =====================================================================

def replace_span_in_paragraph(paragraph, start, end, new_text,
                              color_hex="FF0000"):
    p = paragraph._element
    runs = iter_runs_ordered(p)
    if not runs:
        return False
    full = ''.join(t for _, t in runs)
    if start < 0 or end > len(full) or start >= end:
        return False

    src_rPr = runs[0][0].find(qn('w:rPr'))
    head = full[:start]
    tail = full[end:]

    for ch in list(p):
        if ch.tag in (qn('w:r'), qn('w:hyperlink')):
            p.remove(ch)

    def add_run(text, color=None):
        r = OxmlElement('w:r')
        set_run_plain(r, text, src_rPr=src_rPr, color_hex=color)
        p.append(r)

    if head:
        add_run(head, None)
    add_run(new_text, color_hex)
    if tail:
        add_run(tail, None)
    return True


# =====================================================================
# Клонирование абзаца
# =====================================================================

def clone_paragraph_from(src_p, text, color_hex="FF0000"):
    new_p = copy.deepcopy(src_p)
    for ch in list(new_p):
        if ch.tag in (qn('w:r'), qn('w:hyperlink')):
            new_p.remove(ch)

    pPr = new_p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        new_p.insert(0, pPr)
    jc = pPr.find(qn('w:jc'))
    if jc is not None:
        pPr.remove(jc)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'left')
    pPr.append(jc)

    runs = src_p.findall(qn('w:r'))
    src_rPr = runs[0].find(qn('w:rPr')) if runs else None
    new_r = OxmlElement('w:r')
    set_run_plain(new_r, text, src_rPr=src_rPr, color_hex=color_hex)
    new_p.append(new_r)
    return new_p


def insert_empty_paragraph_like(template_p):
    new_p = copy.deepcopy(template_p)
    for ch in list(new_p):
        if ch.tag in (qn('w:r'), qn('w:hyperlink')):
            new_p.remove(ch)
    runs = template_p.findall(qn('w:r'))
    src_rPr = runs[0].find(qn('w:rPr')) if runs else None
    new_r = OxmlElement('w:r')
    rPr = make_plain_rPr(src_rPr)
    new_r.insert(0, rPr)
    t = OxmlElement('w:t')
    t.text = ''
    t.set(qn('xml:space'), 'preserve')
    new_r.append(t)
    new_p.append(new_r)
    return new_p


def insert_para_after(prev_el, new_p, cell):
    if prev_el is None:
        tcPr = cell._tc.find(qn('w:tcPr'))
        if tcPr is not None:
            tcPr.addnext(new_p)
        else:
            cell._tc.insert(0, new_p)
    else:
        prev_el.addnext(new_p)
    return new_p


# =====================================================================
# Поиск шапки «… , в лице … , именуемое в дальнейшем «Заказчик»»
# =====================================================================

import re as _re

_PREAMBLE_END_RE = _re.compile(
    r',?\s*(?P<imp>именуем\w+)\s+в\s+дальнейшем\s*'
    r'[«"\u201C\u201E]?\s*Заказчик[»"\u201D\u201F]?',
    _re.IGNORECASE)

_NAME_START_RE = _re.compile(
    r'(?<![А-ЯЁа-яё])'
    r'(?:ООО|АО|ПАО|ЗАО|АНО|ИП|ТСН|ТСЖ|ЖСК|ДНП|ДНТ|СНТ|ОНТ|'
    r'МУП|ГУП|ГБУ|МБУ|ГКУ|НКО|СПК|КФХ|ТПП|'
    r'Индивидуальный\s+предприниматель)\s+'
    r'[«"\u201C\u201E]?[А-ЯЁ]',
    _re.IGNORECASE)

_PREAMBLE_CTX_RE = _re.compile(
    r'с\s+одной\s+стороны|с\s+другой\s+стороны|заключили',
    _re.IGNORECASE)


def find_preamble_span(doc):
    for p in iter_all_paragraphs(doc):
        text = p.text or ''
        if not text.strip():
            continue
        if not _PREAMBLE_CTX_RE.search(text):
            continue
        zm = _PREAMBLE_END_RE.search(text)
        if not zm:
            continue
        right_end = zm.end()
        right_before = zm.start()
        while right_before > 0 and text[right_before - 1] in ' ,|':
            right_before -= 1
        imp_start = zm.start('imp')
        imp_end = zm.end('imp')
        imp_word = zm.group('imp')
        name_match = None
        for nm in _NAME_START_RE.finditer(text[:right_before]):
            name_match = nm
        if not name_match:
            continue
        left = name_match.start()
        if right_before <= left:
            continue
        return {
            'p': p,
            'left': left,
            'right_before': right_before,
            'right_end': right_end,
            'imp_start': imp_start,
            'imp_end': imp_end,
            'imp_word': imp_word,
        }
    return None


# =====================================================================
# Поиск блока реквизитов Заказчика
# =====================================================================

_REQ_KEY_RE = _re.compile(
    r'\b(?:ИНН|КПП|ОГРН|ОГРНИП|ГРН|БИК|'
    r'Р\s*/\s*[сcСC]|Расч[её]т|К\s*/\s*[сcСC]|Кор(?:р)?\.?\s*сч|'
    r'Л\s*/\s*[сcСC]|Лицев|'
    r'Юр\.?\s*адрес|Юридический\s+адрес|Почтовый\s+адрес|Адрес)\b',
    _re.IGNORECASE)


def _count_req_keys(text):
    return len(_REQ_KEY_RE.findall(text or ''))


def find_requisites_block(doc):
    best = None
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            cells = iter_unique_cells(row)
            if len(cells) < 2:
                continue
            right = cells[-1]
            text = ' '.join(p.text for p in right.paragraphs)
            marks = _count_req_keys(text)
            if marks < 2:
                continue
            if best is None or marks > best['marks']:
                best = {
                    'kind': 'table',
                    'table': t,
                    'table_idx': ti,
                    'row_idx': ri,
                    'right_cell': right,
                    'marks': marks,
                }
    if best is not None:
        cell = best['right_cell']
        best['paragraphs'] = list(cell.paragraphs)
        return best

    paras = list(iter_all_paragraphs(doc))
    marker_re = _re.compile(r'^\s*(Заказчик)\s*:?\s*$', _re.IGNORECASE)
    for i, p in enumerate(paras):
        if not marker_re.match(p.text or ''):
            continue
        block_paras = []
        for j in range(i + 1, min(i + 60, len(paras))):
            t2 = norm_text(paras[j].text)
            if _re.match(r'^\s*(Подрядчик|Исполнитель|Заказчик)\s*:',
                         t2, _re.I):
                break
            block_paras.append(paras[j])
        if block_paras:
            return {'kind': 'section', 'paragraphs': block_paras}
    return None


_HEAD_LABEL_RE = _re.compile(
    r'^\s*(Заказчик|Исполнитель|Покупатель|Заказчик-2)\s*:?\s*$',
    _re.IGNORECASE)


def replace_requisites_block_in_table(doc, block, new_lines, signature_text):
    cell = block['right_cell']
    old_paras = list(cell.paragraphs)
    if not old_paras:
        return False

    head_paras = []
    body_paras = []
    for p in old_paras:
        t = norm_text(p.text)
        if _HEAD_LABEL_RE.match(t):
            head_paras.append(p)
        else:
            body_paras.append(p)

    if body_paras:
        template_p = body_paras[0]._element
    else:
        template_p = old_paras[0]._element

    for p in body_paras:
        p._element.getparent().remove(p._element)

    prev = head_paras[-1]._element if head_paras else None

    for text in new_lines:
        new_p = clone_paragraph_from(template_p, text)
        prev = insert_para_after(prev, new_p, cell)

    if signature_text:
        empty_p = insert_empty_paragraph_like(template_p)
        prev = insert_para_after(prev, empty_p, cell)
        sig_p = clone_paragraph_from(template_p, signature_text)
        prev = insert_para_after(prev, sig_p, cell)
    return True


def replace_requisites_block_in_section(block, new_lines, signature_text):
    paras = block['paragraphs']
    if not paras:
        return False
    template_p = paras[0]._element
    insert_after = paras[-1]._element
    for text in new_lines:
        new_p = clone_paragraph_from(template_p, text)
        insert_after.addnext(new_p)
        insert_after = new_p
    if signature_text:
        empty_p = insert_empty_paragraph_like(template_p)
        insert_after.addnext(empty_p)
        insert_after = empty_p
        sig_p = clone_paragraph_from(template_p, signature_text)
        insert_after.addnext(sig_p)
        insert_after = sig_p
    for p in paras:
        p._element.getparent().remove(p._element)
    return True


# =====================================================================
# Дата
# =====================================================================

MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня",
          "июля", "августа", "сентября", "октября", "ноября", "декабря"]

DATE_RE = _re.compile(r'«\s*\d{1,2}\s*»\s*[а-яё]+\s*\d{4}\s*г?\.?',
                      _re.IGNORECASE)


def current_date_str():
    t = date.today()
    return f"«{t.day:02d}» {MONTHS[t.month - 1]} {t.year}г."


def _date_is_current(old_date):
    if not old_date:
        return False
    t = date.today()
    has_day = (_re.search(rf'\b{t.day:02d}\b', old_date)
               or _re.search(rf'«\s*0?{t.day}\s*»', old_date))
    has_month = MONTHS[t.month - 1] in old_date.lower()
    has_year = str(t.year) in old_date
    return bool(has_day and has_month and has_year)


def find_first_date(doc):
    for p in iter_all_paragraphs(doc):
        m = DATE_RE.search(p.text or '')
        if m:
            return p, m.start(), m.end(), m.group(0)
    return None


def replace_date_in_doc(doc, color_hex="FF0000"):
    found = find_first_date(doc)
    if not found:
        log.info("Дата не найдена")
        return False
    p, s, e, old_date = found
    if _date_is_current(old_date):
        log.info("Дата уже актуальная: %r", old_date)
        return False
    new_date = current_date_str()
    ok = replace_span_in_paragraph(p, s, e, new_date, color_hex=color_hex)
    log.info("Дата: %r -> %r (%s)", old_date, new_date, ok)
    return ok
