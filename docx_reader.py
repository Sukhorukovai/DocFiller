# -*- coding: utf-8 -*-
"""
Чтение .docx в текст.

Три способа, в порядке приоритета:
  1. python-docx (обычные файлы).
  2. Прямой XML-обход (strict OOXML / SmartArt / текстовые поля).
  3. Word → PDF → OCR (для «немых» файлов, где ничего не читается).

Плюс OCR картинок из word/media/.
"""
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from ..logging_setup import log
from ..core.text import norm_text, is_zip
from . import ocr as _ocr
from .word_bridge import (
    doc_to_docx_via_word, docx_resave_via_word, docx_to_pdf_via_word,
)

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import cv2
    import numpy as np
    from PIL import Image
except ImportError:
    cv2 = None


# =====================================================================
# Ячейки таблиц
# =====================================================================

_W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _para_text_deep(p) -> str:
    """
    Как p.text, но проходит вглубь — внутрь w:smartTag / w:sdt /
    w:hyperlink. python-docx::Paragraph.text собирает только
    прямых детей <w:p>/<w:r>, поэтому текст, упакованный Word'ом
    в смарт-тег (например, «metricconverter», который Word вешает
    на почтовый индекс вида «603043, г»), молча теряется.
    """
    pieces = []
    for node in p._p.iter():
        tag = node.tag
        if not isinstance(tag, str):
            continue
        if tag == _W_NS + 't':
            if node.text:
                pieces.append(node.text)
        elif tag == _W_NS + 'tab':
            pieces.append('\t')
        elif tag in (_W_NS + 'br', _W_NS + 'cr'):
            pieces.append('\n')
    return ''.join(pieces)


def cell_text(cell) -> str:
    parts = [_para_text_deep(p) for p in cell.paragraphs]
    for t in cell.tables:
        for row in t.rows:
            parts.append(' | '.join(c.text.strip() for c in row.cells
                                    if c.text.strip()))
    return '\n'.join(p for p in parts if p)


def iter_unique_cells(row) -> list:
    """Убирает объединённые ячейки (одна tc может повторяться в row.cells)."""
    seen = set()
    result = []
    for c in row.cells:
        tc = c._tc
        if id(tc) in seen:
            continue
        seen.add(id(tc))
        result.append(c)
    return result


# =====================================================================
# Нормализация текста из docx
# =====================================================================

def join_split_addr_labels(text: str) -> str:
    """«Юридический» + «адрес: …» → «Юридический адрес: …»."""
    lines = text.split('\n')
    out = []
    i = 0
    while i < len(lines):
        cur = lines[i].rstrip()
        m = re.match(r'^\s*(Юридический|Почтовый|Фактический|Юр\.?)\s*$',
                     cur, re.I)
        if m and i + 1 < len(lines):
            nxt = lines[i + 1].lstrip()
            if re.match(r'^адрес\b', nxt, re.I):
                out.append(cur.strip() + ' ' + nxt)
                i += 2
                continue
        m2 = re.match(
            r'^\s*(Юридический\s+адрес|Почтовый\s+адрес|'
            r'Фактический\s+адрес|Юр\.\s*адрес)\s*$',
            cur, re.I)
        if m2 and i + 1 < len(lines):
            nxt = lines[i + 1].lstrip()
            if nxt.startswith(':'):
                out.append(cur.strip() + nxt)
                i += 2
                continue
        out.append(cur)
        i += 1
    return '\n'.join(out)


def join_broken_labels(text: str) -> str:
    """Склеивает разорванные метки (Банк / Наименование банка / Расчётный счёт)."""
    lines = text.split('\n')
    out = []
    i = 0
    broken = re.compile(
        r'^\s*(Местонахождение|Юридический\s+адрес|Почтовый\s+адрес|'
        r'Фактический\s+адрес|Место\s+жительства|Банк|'
        r'Наименование\s+банка|Полное\s+наименование\s+банка|'
        r'Расчетный\s+счет|Лицевой\s+счет|Корреспондентский\s+счет)\s*$',
        re.I)
    while i < len(lines):
        cur = lines[i].rstrip()
        if broken.match(cur):
            merged = cur.strip()
            j = i + 1
            while j < len(lines) and j < i + 6:
                nxt = lines[j].strip()
                if not nxt:
                    j += 1
                    continue
                merged += ' ' + nxt
                j += 1
                if ':' in merged or len(merged) > 300:
                    break
            out.append(merged)
            i = j
        else:
            out.append(cur)
            i += 1
    return '\n'.join(out)


_TABLE_LABEL_HINTS = re.compile(
    r'(ИНН|КПП|ОГРН|ОГРНИП|ГРН|БИК|Р/с|Р/С|Р/сч|К/с|К/С|Корр|Кор\.|'
    r'Корреспондентск|Л/с|Л/С|Лицев|Расч[её]тн|расчетный\s+счет|'
    r'№\s*сч[её]та|адрес|Адрес|нахожден|счет|счёт|банк|Банк|'
    r'Телефон|Факс|E-mail|Эл\.\s*почта|Электронная\s+почта|'
    r'[Дд]иректор|правления|компании|наименование|Наименование|'
    r'Сокращенно|Владелец|Руководитель|Управляющий|'
    r'ИМК|Соглашение|ОКПО|ОКВЭД|ОКТМО)',
    re.IGNORECASE)


def normalize_table_lines(text: str) -> str:
    """«Метка | Значение» → «Метка: Значение» (для строк из таблиц docx)."""
    if ' | ' not in text:
        return text
    out = []
    for line in text.split('\n'):
        if ' | ' not in line:
            out.append(line)
            continue
        parts = [p.strip() for p in line.split(' | ')]
        parts = [p for p in parts if p]
        if not parts:
            continue
        result_parts = []
        i = 0
        while i < len(parts):
            cur = parts[i]
            if _TABLE_LABEL_HINTS.search(cur) and i + 1 < len(parts):
                lbl = re.sub(r'\s*[:\-]\s*$', '', cur)
                if _TABLE_LABEL_HINTS.search(parts[i + 1]) and ':' in parts[i + 1]:
                    result_parts.append(f"{lbl}:")
                    i += 1
                else:
                    result_parts.append(f"{lbl}: {parts[i + 1]}")
                    i += 2
            else:
                result_parts.append(cur)
                i += 1
        out.append('\n'.join(result_parts))
    return '\n'.join(out)


# =====================================================================
# XML-обход (strict OOXML и «немые» файлы)
# =====================================================================

def docx_text_from_xml(path: Path) -> str:
    """
    Достаёт текст из word/document.xml, определяя namespace из корня.
    Ловит strict OOXML (где python-docx не открывает).
    """
    default_ns = ('http://schemas.openxmlformats.org/'
                  'wordprocessingml/2006/main')
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read('word/document.xml')
    except Exception as e:
        log.info("docx_text_from_xml: не читается %s: %s", path.name, e)
        return ""

    try:
        root = ET.fromstring(xml)
    except Exception as e:
        log.info("docx_text_from_xml: битый XML %s: %s", path.name, e)
        return ""

    m = re.match(r'\{([^}]+)\}', root.tag)
    ns_url = m.group(1) if m else default_ns
    log.info("docx_text_from_xml: namespace = %s", ns_url)

    ns = '{' + ns_url + '}'
    lines = []
    for p in root.iter(ns + 'p'):
        t = ''.join(el.text or '' for el in p.iter(ns + 't'))
        if t.strip():
            lines.append(t)
    if lines:
        return "\n".join(lines)

    # Fallback без namespace — на экзотические случаи.
    lines = []
    for p in root.iter():
        tag = p.tag
        if not isinstance(tag, str):
            continue
        if not tag.endswith('}p') and tag != 'p':
            continue
        pieces = []
        for el in p.iter():
            t2 = el.tag
            if not isinstance(t2, str):
                continue
            if t2.endswith('}t') or t2 == 't':
                if el.text:
                    pieces.append(el.text)
        t = ''.join(pieces)
        if t.strip():
            lines.append(t)
    return "\n".join(lines)


def docx_ocr_media(path: Path) -> str:
    """OCR всех картинок из word/media/ (если есть)."""
    if _ocr.pytesseract is None or cv2 is None:
        return ""
    try:
        with zipfile.ZipFile(path) as z:
            imgs = [n for n in z.namelist()
                    if n.startswith('word/media/')
                    and re.search(r'\.(png|jpg|jpeg|bmp|tif|tiff)$',
                                  n, re.I)]
        if not imgs:
            log.info("docx_ocr_media(%s): в word/media/ нет картинок",
                     path.name)
            return ""
        ocr_chunks = []
        for name in imgs:
            data = z.read(name)
            arr = np.frombuffer(data, dtype=np.uint8)
            img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
            if img is None:
                try:
                    pil = Image.open(io.BytesIO(data)).convert("RGB")
                    img = cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)
                except Exception:
                    img = None
            if img is not None and img.size > 0:
                ocr_chunks.append(_ocr.ocr_image(img))
        return "\n".join(ocr_chunks)
    except Exception as e:
        log.warning("docx_ocr_media(%s): %s", path.name, e)
        return ""


def from_docx_via_pdf_ocr(path: Path) -> str:
    """Word → PDF → _from_pdf. Требует pywin32 и хотя бы pdfplumber."""
    from .pdf_reader import from_pdf  # локальный импорт, чтобы не было цикла

    from ..logging_setup import log as _log
    try:
        import win32com  # noqa: F401
    except ImportError:
        _log.info("from_docx_via_pdf_ocr: win32com отсутствует — пропуск")
        return ""

    try:
        pdf_path = docx_to_pdf_via_word(path)
    except Exception as e:
        _log.warning("Word→PDF не удалось для %s: %s", path.name, e)
        return ""
    try:
        text = from_pdf(pdf_path)
    except Exception as e:
        _log.warning("PDF→текст не удалось для %s: %s", pdf_path.name, e)
        return ""
    _log.info("from_docx_via_pdf_ocr: получено %d символов", len(text or ""))
    return text or ""


# =====================================================================
# Главная функция
# =====================================================================

def from_docx(path: Path, _resave: bool = True) -> str:
    if not is_zip(path):
        try:
            conv = doc_to_docx_via_word(path)
            return from_docx(conv)
        except Exception as e:
            log.warning("doc→docx не удалось для %s: %s", path.name, e)

    parts = []
    doc_err = None
    try:
        doc = Document(path)
        for p in doc.paragraphs:
            parts.append(_para_text_deep(p))

        def walk_tables(tables):
            for t in tables:
                for row in t.rows:
                    cells_text = []
                    for c in iter_unique_cells(row):
                        txt = cell_text(c).strip()
                        if txt:
                            cells_text.append(txt)
                    line = ' | '.join(cells_text)
                    if line:
                        parts.append(line)
                    for c in iter_unique_cells(row):
                        walk_tables(c.tables)

        walk_tables(doc.tables)
    except Exception as e:
        doc_err = e
        log.info("Document() упал на %s: %s", path.name, e)

    text = "\n".join(parts)
    if len(text.strip()) < 5:
        text = docx_text_from_xml(path)
        if text:
            log.info("from_docx(%s): XML-обход дал %d символов",
                     path.name, len(text))

    if len(text.strip()) < 20:
        ocr_media = docx_ocr_media(path)
        if ocr_media:
            text = (text + "\n" + ocr_media).strip()
            log.info("from_docx(%s): после OCR картинок %d символов",
                     path.name, len(text))

    if len(text.strip()) < 20:
        log.info("from_docx(%s): текста нет — пробую Word→PDF→OCR",
                 path.name)
        pdf_ocr_text = from_docx_via_pdf_ocr(path)
        if pdf_ocr_text and len(pdf_ocr_text.strip()) > len(text.strip()):
            text = pdf_ocr_text

    if (len(text.strip()) < 20 and _resave
            and _win32_available()):
        log.info("from_docx(%s): пробую пересохранить через Word",
                 path.name)
        try:
            new_docx = docx_resave_via_word(path)
            if (new_docx and new_docx.exists()
                    and new_docx.resolve() != path.resolve()):
                text2 = from_docx(new_docx, _resave=False)
                if len(text2.strip()) > len(text.strip()):
                    text = text2
                    log.info("from_docx(%s): пересохранение помогло, %d",
                             path.name, len(text))
        except Exception as e:
            log.info("from_docx(%s): Word-resave не помог: %s",
                     path.name, e)

    if len(text.strip()) < 20:
        log.warning("from_docx(%s): НЕ УДАЛОСЬ извлечь текст. "
                    "doc_err=%r, есть_ли_word=%s",
                    path.name, doc_err, _win32_available())

    text = join_split_addr_labels(text)
    text = join_broken_labels(text)
    text = normalize_table_lines(text)
    return text


def _win32_available() -> bool:
    try:
        import win32com  # noqa: F401
        return True
    except ImportError:
        return False